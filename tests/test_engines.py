from __future__ import annotations

import os
import subprocess
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from docmorph.core.models import Format, Mode, PdfKind, Route
from docmorph.engines.images import ImageToPdfEngine
from docmorph.engines.libreoffice import LibreOfficeEngine
from docmorph.engines.ocr import OcrTextEngine
from docmorph.engines.pdf_layout import PdfLayoutEngine
from docmorph.engines.pdf_structure import PdfStructureEngine
from docmorph.engines.pdf_text import PdfTextEngine
from tests.conftest import LEFT_TEXT, RIGHT_TEXT, requires


def _skip_unless_available(engine) -> None:
    reason = engine.availability()
    if reason:
        pytest.skip(reason)


def _pdf_text(path: Path) -> str:
    with pymupdf.open(path) as doc:
        return "".join(page.get_text() for page in doc)


def test_pdf_text_reads_columns_in_order(two_column_pdf: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.txt"
    PdfTextEngine().convert(two_column_pdf, out, Route(Format.PDF, Format.TXT, pdf_kind=PdfKind.DIGITAL))
    text = out.read_text(encoding="utf-8")
    title = text.index("Titre")
    left = text.index(LEFT_TEXT.split(".")[0])
    right = text.index(RIGHT_TEXT.split(".")[0])
    assert title < left < right


@pytest.mark.parametrize("fixture", ["png_with_alpha", "jpg_file"])
def test_image_to_pdf(fixture: str, request: pytest.FixtureRequest, tmp_path: Path) -> None:
    src: Path = request.getfixturevalue(fixture)
    out = tmp_path / "image.pdf"
    fmt = Format.PNG if src.suffix == ".png" else Format.JPG
    ImageToPdfEngine().convert(src, out, Route(fmt, Format.PDF))
    with pymupdf.open(out) as doc:
        assert doc.page_count == 1
        assert doc[0].get_images()


@pytest.mark.libreoffice
@requires("soffice")
def test_libreoffice_docx_to_pdf(docx_file: Path, tmp_path: Path) -> None:
    out = tmp_path / "rapport.pdf"
    LibreOfficeEngine().convert(docx_file, out, Route(Format.DOCX, Format.PDF))
    assert "Rapport DocMorph" in _pdf_text(out)


@pytest.mark.libreoffice
@requires("soffice")
def test_libreoffice_text_encodings(txt_utf8: Path, txt_cp1252: Path, tmp_path: Path) -> None:
    out_utf8 = tmp_path / "utf8.pdf"
    out_cp = tmp_path / "cp.pdf"
    LibreOfficeEngine().convert(txt_utf8, out_utf8, Route(Format.TXT, Format.PDF))
    LibreOfficeEngine().convert(txt_cp1252, out_cp, Route(Format.TXT, Format.PDF))
    assert "accentuée" in _pdf_text(out_utf8)
    assert "été" in _pdf_text(out_cp)


@pytest.mark.libreoffice
@requires("soffice")
def test_libreoffice_rtf_to_docx_and_doc(rtf_file: Path, tmp_path: Path) -> None:
    docx = tmp_path / "lettre.docx"
    LibreOfficeEngine().convert(rtf_file, docx, Route(Format.RTF, Format.DOCX))
    assert "Bonjour RTF" in "\n".join(p.text for p in Document(docx).paragraphs)
    doc = tmp_path / "lettre.doc"
    LibreOfficeEngine().convert(docx, doc, Route(Format.DOCX, Format.DOC))
    assert doc.read_bytes().startswith(bytes.fromhex("D0CF11E0A1B11AE1"))


def test_pdf_layout_to_docx(two_column_pdf: Path, tmp_path: Path) -> None:
    engine = PdfLayoutEngine()
    _skip_unless_available(engine)
    out = tmp_path / "article.docx"
    engine.convert(two_column_pdf, out, Route(Format.PDF, Format.DOCX, Mode.LAYOUT, PdfKind.DIGITAL))
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Titre de l'article" in text


@pytest.mark.ocr
def test_ocr_scanned_pdf_to_text(scanned_pdf: Path, tmp_path: Path) -> None:
    engine = OcrTextEngine()
    _skip_unless_available(engine)
    out = tmp_path / "scan.txt"
    engine.convert(scanned_pdf, out, Route(Format.PDF, Format.TXT, pdf_kind=PdfKind.SCANNED))
    assert "DOCMORPH" in out.read_text(encoding="utf-8").upper()


_LATEX_ARTICLE = r"""
\documentclass[twocolumn]{article}
\usepackage{amsmath}
\begin{document}
\section{Introduction}
We study the heat equation on a bounded domain and recall the classical energy estimate.
The solution satisfies the following identity for every time $t > 0$:
\begin{equation}
  \frac{d}{dt} \int_{\Omega} u^2 \, dx = -2 \int_{\Omega} |\nabla u|^2 \, dx .
\end{equation}
Integrating in time gives a bound that depends only on the initial data.
\begin{equation}
  \sum_{k=1}^{n} k^2 = \frac{n(n+1)(2n+1)}{6}
\end{equation}
\section{Conclusion}
These identities are standard and serve as a test of formula recognition.
\end{document}
"""


@pytest.mark.structure
@pytest.mark.skipif(not os.environ.get("DOCMORPH_TEST_STRUCTURE"), reason="set DOCMORPH_TEST_STRUCTURE=1 (slow)")
@requires("pdflatex")
def test_structure_mode_produces_word_equations(tmp_path: Path) -> None:
    engine = PdfStructureEngine()
    _skip_unless_available(engine)
    (tmp_path / "article.tex").write_text(_LATEX_ARTICLE, encoding="utf-8")
    subprocess.run(
        ["pdflatex", "-interaction=nonstopmode", "article.tex"], cwd=tmp_path, check=True, capture_output=True
    )
    out = tmp_path / "article.docx"
    engine.convert(tmp_path / "article.pdf", out, Route(Format.PDF, Format.DOCX, Mode.STRUCTURE, PdfKind.DIGITAL))

    assert "Introduction" in "\n".join(p.text for p in Document(out).paragraphs)
    with zipfile.ZipFile(out) as zf:
        assert "<m:oMath" in zf.read("word/document.xml").decode("utf-8")
