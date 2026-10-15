from __future__ import annotations

import io
import shutil
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from PIL import Image, ImageDraw

LEFT_TEXT = "Colonne gauche premier paragraphe. " * 6
RIGHT_TEXT = "Colonne droite second paragraphe. " * 6


def requires(binary: str):
    return pytest.mark.skipif(shutil.which(binary) is None, reason=f"{binary} not installed")


@pytest.fixture
def two_column_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "article.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_textbox(pymupdf.Rect(50, 40, 545, 80), "Titre de l'article sur toute la largeur", fontsize=14)
    page.insert_textbox(pymupdf.Rect(50, 100, 290, 400), LEFT_TEXT, fontsize=10)
    page.insert_textbox(pymupdf.Rect(305, 100, 545, 400), RIGHT_TEXT, fontsize=10)
    doc.save(path)
    return path


@pytest.fixture
def scanned_pdf(tmp_path: Path) -> Path:
    img = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(["DOCMORPH OCR TEST", "Bonjour le monde", "Hello world 2026"]):
        draw.text((120, 150 + i * 120), line, fill="black", font_size=72)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    path = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(path)
    return path


@pytest.fixture
def png_with_alpha(tmp_path: Path) -> Path:
    path = tmp_path / "logo.png"
    Image.new("RGBA", (200, 100), (255, 0, 0, 128)).save(path)
    return path


@pytest.fixture
def jpg_file(tmp_path: Path) -> Path:
    path = tmp_path / "photo.jpg"
    Image.new("RGB", (320, 200), "navy").save(path, quality=90)
    return path


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    path = tmp_path / "rapport.docx"
    document = Document()
    document.add_heading("Rapport DocMorph", 0)
    document.add_paragraph("Paragraphe de test avec accents : éèàçù.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(1, 1).text = "D"
    document.save(path)
    return path


@pytest.fixture
def txt_utf8(tmp_path: Path) -> Path:
    path = tmp_path / "notes.txt"
    path.write_text("Ligne un — accentuée é\nسطر عربي\n", encoding="utf-8")
    return path


@pytest.fixture
def txt_cp1252(tmp_path: Path) -> Path:
    path = tmp_path / "ancien.txt"
    path.write_bytes("Texte ancien encodé en Windows-1252 : été\n".encode("cp1252"))
    return path


@pytest.fixture
def rtf_file(tmp_path: Path) -> Path:
    path = tmp_path / "lettre.rtf"
    path.write_text(r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Times;}}\f0\fs24 Bonjour RTF \b gras\b0 .\par}", encoding="ascii")
    return path


def make_zip(path: Path, entries: dict[str, str], stored_first: str | None = None) -> Path:
    with zipfile.ZipFile(path, "w") as zf:
        if stored_first:
            zf.writestr(stored_first, entries.pop(stored_first), compress_type=zipfile.ZIP_STORED)
        for name, content in entries.items():
            zf.writestr(name, content)
    return path
